"""Rebuild the model observation documents and comparison tables from saved results.

No training, source data, GPU, or pickle loading is required. Install
``requirements-docs.txt`` and run this file from any working directory. Source
tables are read only. Both document formats are generated from the same content.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import io
import json
import statistics
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


MODELS = ("logistic_regression", "tabm")
NAMES = {"logistic_regression": "Logistic regression", "tabm": "TabM"}
ARMS = ("A", "B", "C", "D")
POLICY = "fixed_0.50"
SCENARIOS = ("observed", "missing_1", "missing_3", "missing_6", "noise_0.05",
             "noise_0.10", "noise_0.20", "gain_0.05", "gain_0.10", "gain_0.20",
             "combined_3lead_noise0.10")
LABELS = {"observed": "Observed", "missing_1": "1 missing lead",
          "missing_3": "3 missing leads", "missing_6": "6 missing leads",
          "combined_3lead_noise0.10": "3 leads + noise 0.10"}
LABELS.update({f"{kind}_{dose}": f"{kind.title()} {dose}"
               for kind in ("noise", "gain") for dose in ("0.05", "0.10", "0.20")})


class Results:
    def __init__(self, root: Path):
        self.root, self.inputs, self.cache = root, {}, {}
        self.config = {m: self.json(f"results/{m}/run_config.json") for m in MODELS}
        self.complete = {m: self.json(f"results/{m}/reports/completion.json") for m in MODELS}
        if not self.complete["tabm"].get("complete_tabm_plan"):
            raise ValueError("This observation report requires a complete TabM run.")
        if not self.complete["logistic_regression"].get("complete_logistic_plan"):
            raise ValueError("This observation report requires a complete logistic run.")
        if self.config[MODELS[0]]["manifest"] != self.config[MODELS[1]]["manifest"]:
            raise ValueError("The two runs must use the same prepared data manifest.")
        for m in MODELS:
            rows = self.csv(m, "summary")
            self.cache[m, "summary_index"] = {
                (r["variant"], r["scenario"], r["threshold_policy"], r["metric"]): r
                for r in rows}
            if len(self.cache[m, "summary_index"]) != len(rows):
                raise ValueError(f"Duplicate summary keys for {m}.")
            for arm in ARMS:
                for scenario in SCENARIOS:
                    for metric in ("macro_f1", "sensitivity", "specificity", "brier"):
                        r = self.row(m, arm, metric, scenario)
                        if int(r["n_folds"]) != 10 or int(r["n_defined"]) != 10:
                            raise ValueError(f"Incomplete metric: {m}/{arm}/{scenario}/{metric}")
        self.pairing = self.verify_stress_pairing()

    def verify_stress_pairing(self):
        """Check common fold-variant stress audits without loading model checkpoints."""
        variants = sorted(set(self.config[MODELS[0]]["variants"]) &
                          set(self.config[MODELS[1]]["variants"]))
        conditions = 0
        for outer in range(10):
            for variant in variants:
                audits = []
                for model in MODELS:
                    relative = (f"results/{model}/output_data/folds/{outer:02d}_{variant}/"
                                "perturbation_audit.json")
                    rows = self.json(relative)["test"]
                    if len(rows) != 31:
                        raise ValueError(f"Expected 31 stress conditions in {relative}.")
                    audits.append([{k: v for k, v in row.items() if k != "model"} for row in rows])
                if audits[0] != audits[1]:
                    raise ValueError(f"Cross-model perturbation audit mismatch: {outer}/{variant}.")
                conditions += len(audits[0])
        return {"status": "passed", "shared_prepared_manifest": True,
                "fold_variants_compared": 10 * len(variants), "conditions_compared": conditions,
                "excluded_metadata_fields": ["model"],
                "scope": "Exact equality of saved test perturbation diagnostics; no models refitted."}

    def read(self, relative):
        payload = (self.root / relative).read_bytes()
        self.inputs[relative] = {"sha256": hashlib.sha256(payload).hexdigest(), "bytes": len(payload)}
        return payload.decode("utf-8-sig")

    def json(self, relative):
        return json.loads(self.read(relative))

    def csv(self, model, stem):
        key = (model, stem)
        if key not in self.cache:
            self.cache[key] = list(csv.DictReader(io.StringIO(
                self.read(f"results/{model}/output_data/{stem}.csv"))))
        return self.cache[key]

    def row(self, model, arm, metric, scenario="observed", policy=POLICY):
        return self.cache[model, "summary_index"][arm, scenario, policy, metric]

    def value(self, model, arm, metric="macro_f1", scenario="observed", policy=POLICY):
        return float(self.row(model, arm, metric, scenario, policy)["mean"])

    def fmt(self, model, arm, metric="macro_f1", scenario="observed", policy=POLICY, sd=False):
        r = self.row(model, arm, metric, scenario, policy)
        value = f"{float(r['mean']):.4f}"
        return value + (f" ({float(r['sd']):.4f})" if sd else "")

    def difference(self, model, comparison, scenario="observed", metric="macro_f1"):
        rows = [r for r in self.csv(model, "paired_method_differences")
                if r["comparison"] == comparison and r["scenario"] == scenario
                and r["metric"] == metric and r.get("threshold_policy", POLICY) == POLICY]
        if len(rows) != 1:
            raise ValueError(f"Expected one paired difference: {model}, {comparison}, {scenario}")
        return float(rows[0]["mean_difference"]), float(rows[0]["sd_difference"])


class Report:
    """Build matching Word and Markdown with modest, readable report styling."""

    def __init__(self):
        self.doc, self.md = Document(), []
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.9)
        section.left_margin = section.right_margin = Cm(1.8)
        for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style.font.color.rgb = RGBColor(0, 0, 0)
        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.08
        self.doc.styles["Title"].font.size = Pt(24)
        self.doc.styles["Heading 1"].font.size = Pt(16)
        self.doc.styles["Heading 2"].font.size = Pt(12)
        for name in ("Heading 1", "Heading 2"):
            self.doc.styles[name].paragraph_format.space_before = Pt(12)
            self.doc.styles[name].paragraph_format.space_after = Pt(6)
        self.doc.core_properties.title = "Model observations"
        self.doc.core_properties.author = "COMP90051 project team"
        self.doc.core_properties.subject = "Results, limitations and further evaluation"
        stamp = datetime(2026, 10, 15, tzinfo=timezone.utc)
        self.doc.core_properties.created = self.doc.core_properties.modified = stamp

    def title(self, text):
        self.doc.add_paragraph(text, "Title")
        self.md.append(f"# {text}\n")

    def heading(self, text, level=1):
        self.doc.add_heading(text, level)
        self.md.append(f"{'#' * (level + 1)} {text}\n")

    def p(self, text):
        self.doc.add_paragraph(text)
        self.md.append(text + "\n")

    def code(self, lines):
        paragraph = self.doc.add_paragraph()
        for index, line in enumerate(lines):
            run = paragraph.add_run(("\n" if index else "") + line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        self.md.extend(["```powershell", *lines, "```", ""])

    def page(self):
        self.doc.add_page_break()

    def table(self, headers, rows, widths=None):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = False
        if widths is None:
            widths = [17.4 / len(headers)] * len(headers)
        for c, width in zip(table.columns, widths):
            c.width = Cm(width)
        for cell, width, text in zip(table.rows[0].cells, widths, headers):
            cell.width = Cm(width)
            cell.text = str(text)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EEF3")
            cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for values in rows:
            for cell, width, text in zip(table.add_row().cells, widths, values):
                cell.width = Cm(width)
                cell.text = str(text)
        for row in table.rows:
            prevent_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(prevent_split)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.space_before = Pt(4)
                    paragraph.paragraph_format.line_spacing = 1
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
        self.md.append("| " + " | ".join(headers) + " |")
        self.md.append("| " + " | ".join("---" for _ in headers) + " |")
        self.md.extend("| " + " | ".join(html.escape(str(x)).replace("\n", " ").replace("|", "\\|")
                                            for x in row) + " |" for row in rows)
        self.md.append("")

    def save(self, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        stream = io.BytesIO()
        self.doc.save(stream)
        # Normalize ZIP member timestamps so identical inputs produce identical files.
        with ZipFile(stream) as source, ZipFile(path, "w", ZIP_DEFLATED) as dest:
            for member in source.infolist():
                info = ZipInfo(member.filename, (2026, 10, 15, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                dest.writestr(info, source.read(member.filename))
        path.with_suffix(".md").write_text("\n".join(self.md), encoding="utf-8", newline="\n")

