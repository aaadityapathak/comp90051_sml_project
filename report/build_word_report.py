"""Build an editable Word report from the same content used for the PDF.

Run after report/analyze_results.py. Importing build_report reads the saved
comparison tables but does not rebuild or overwrite the submission PDF.
"""

from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Mm, Pt, RGBColor

from build_report import content_blocks


REPORT_DIRECTORY = Path(__file__).resolve().parent
OUTPUT_DOCX = REPORT_DIRECTORY / 'COMP90051_Final_Report.docx'
FONT = 'Times New Roman'


class ParagraphWriter(HTMLParser):
    """Translate the PDF's inline emphasis and links into editable Word runs."""

    def __init__(self, paragraph):
        super().__init__(convert_charrefs=True)
        self.paragraph = paragraph
        self.bold = False
        self.italic = False
        self.hyperlink = None

    def handle_starttag(self, tag, attributes):
        if tag == 'b':
            self.bold = True
        elif tag == 'i':
            self.italic = True
        elif tag == 'br':
            self.paragraph.add_run().add_break()
        elif tag == 'link':
            target = dict(attributes)['href']
            relation = self.paragraph.part.relate_to(
                target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            self.hyperlink = OxmlElement('w:hyperlink')
            self.hyperlink.set(qn('r:id'), relation)
            self.paragraph._p.append(self.hyperlink)

    def handle_endtag(self, tag):
        if tag == 'b':
            self.bold = False
        elif tag == 'i':
            self.italic = False
        elif tag == 'link':
            self.hyperlink = None

    def handle_data(self, text):
        run = self.paragraph.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if self.bold:
            run.bold = True
        if self.italic:
            run.italic = True
        if self.hyperlink is not None:
            self.hyperlink.append(run._r)


def set_styles(document):
    """Match the report's restrained, single-column academic typography."""
    normal = document.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = Pt(11.6)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name in ('Title', 'Heading 1', 'Caption'):
        style = document.styles[name]
        style.font.name = FONT
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name in ('Title', 'Heading 1')
        style.font.italic = False
        style.paragraph_format.line_spacing = Pt(11.6)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.space_before = Pt(
            6 if name == 'Heading 1' else 3 if name == 'Caption' else 0)
        style.paragraph_format.keep_with_next = name in ('Title', 'Heading 1')
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)


def add_table(document, block):
    """Keep numeric cells editable and repeat headers if a table spans pages."""
    table = document.add_table(rows=1, cols=len(block['headers']))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    properties = table._tbl.tblPr
    # Word aligns cell text to the margin by default. Include the left padding
    # in the table indent so its outer border also stays within the margins.
    indent = OxmlElement('w:tblInd')
    indent.set(qn('w:w'), '100')
    indent.set(qn('w:type'), 'dxa')
    properties.append(indent)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + side)
        for key, value in (('val', 'single'), ('sz', '3'), ('color', 'D9D9D9')):
            border.set(qn('w:' + key), value)
        borders.append(border)
    properties.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, twips in (('top', 60), ('bottom', 60), ('left', 100), ('right', 100)):
        margin = OxmlElement('w:' + side)
        margin.set(qn('w:w'), str(twips))
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    properties.append(margins)
    for index, width in enumerate(block['widths']):
        table.columns[index].width = Mm(178 * width)
    for row_index, values in enumerate([block['headers']] + block['data']):
        row = table.rows[0] if row_index == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for column, value in enumerate(values):
            cell = row.cells[column]
            cell.width = Mm(178 * block['widths'][column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(11.2)
            paragraph.paragraph_format.keep_with_next = row_index == 0
            run = paragraph.add_run(str(value))
            run.font.name = FONT
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.bold = row_index == 0
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E9EEF2')
                cell._tc.get_or_add_tcPr().append(shade)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)
    spacer.add_run().font.size = Pt(1)


def main():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(16)
    section.top_margin = section.bottom_margin = Mm(16)
    section.footer_distance = Mm(16)
    set_styles(document)
    pending_page_break = False
    for block in content_blocks:
        kind = block['type']
        if kind == 'page_break':
            pending_page_break = True
        elif kind == 'paragraph':
            style = {'title': 'Title', 'heading': 'Heading 1',
                     'caption': 'Caption'}.get(block['style'], 'Normal')
            paragraph = document.add_paragraph(style=style)
            if pending_page_break:
                paragraph.paragraph_format.page_break_before = True
                pending_page_break = False
            ParagraphWriter(paragraph).feed(block['text'])
            if block['style'] == 'reference':
                paragraph.paragraph_format.left_indent = Pt(14)
                paragraph.paragraph_format.first_line_indent = Pt(-14)
                paragraph.paragraph_format.line_spacing = Pt(11.1)
                paragraph.paragraph_format.space_after = Pt(4)
            if block['style'] == 'caption':
                paragraph.paragraph_format.line_spacing = Pt(11.2)
                paragraph.paragraph_format.keep_with_next = block['text'].startswith('<b>Table')
        elif kind == 'table':
            add_table(document, block)
        elif kind == 'image':
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.keep_with_next = True
            shape = paragraph.add_run().add_picture(
                str(REPORT_DIRECTORY / block['path']),
                width=Mm(block['width_mm']), height=Mm(block['height_mm']))
            shape._inline.docPr.set(
                'descr', 'Paired augmentation effects on macro-F1 for LR, HGB and TabM. '
                'Panels show observed measurements and six hidden leads; error bars '
                'are sample SD across ten folds. The panels use different scales.')

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    footer._p.append(page)
    document.core_properties.title = 'AFIB classification under missing and perturbed ECG measurements'
    document.core_properties.author = 'Jiayi Huang; Siwen Wang; Aaditya Pathak'
    document.core_properties.subject = 'COMP90051 Statistical Machine Learning final report'
    document.core_properties.comments = ''
    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == '__main__':
    main()
